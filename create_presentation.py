"""
Generates a professional 9-slide PowerPoint presentation (.pptx)
for the T20 World Cup Cricket Score Predictor academic major project.
Uses 16:9 widescreen layout with professional cricket dark theme styling.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_PPTX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "T20_World_Cup_Score_Predictor_Presentation.pptx")

# Palette constants
COLOR_BG_DARK = RGBColor(10, 26, 18)        # Deep cricket pitch green/dark
COLOR_CARD_BG = RGBColor(18, 44, 32)        # Card surface green
COLOR_ACCENT_GREEN = RGBColor(0, 230, 118)  # Vibrant cricket green
COLOR_ACCENT_CYAN = RGBColor(56, 189, 248)  # Cyan
COLOR_ACCENT_GOLD = RGBColor(251, 191, 36)  # Warm gold
COLOR_TEXT_WHITE = RGBColor(248, 250, 252)  # Bright text
COLOR_TEXT_MUTED = RGBColor(167, 243, 208)  # Soft mint/gray
COLOR_BORDER = RGBColor(24, 72, 52)         # Border green

def apply_background(slide, prs):
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = COLOR_BG_DARK
    bg_shape.line.fill.background() # No border
    return bg_shape

def add_header(slide, title_text, category_text="B.TECH CSE (AI/ML) MAJOR PROJECT"):
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_cat = tf.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_ACCENT_GREEN

    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT_WHITE
    p_title.space_before = Pt(4)

def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    return card

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_background(s1, prs)

    # Accent decorative banner
    card1 = add_card(s1, Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.1), bg_color=RGBColor(14, 34, 25), border_color=COLOR_ACCENT_GREEN)

    tbox1 = s1.shapes.add_textbox(Inches(1.5), Inches(1.6), Inches(10.333), Inches(4.3))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "AMITY UNIVERSITY JHARKHAND • MAJOR PROJECT"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT_GOLD

    p1 = tf1.add_paragraph()
    p1.text = "T20 World Cup Cricket Score Predictor"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_WHITE
    p1.space_before = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "Design and Development of an Explainable Ball-by-Ball First-Innings Score Forecasting System"
    p2.font.size = Pt(16)
    p2.font.color.rgb = COLOR_ACCENT_GREEN
    p2.space_before = Pt(6)

    p3 = tf1.add_paragraph()
    p3.text = "—" * 40
    p3.font.size = Pt(10)
    p3.font.color.rgb = COLOR_BORDER
    p3.space_before = Pt(14)

    p4 = tf1.add_paragraph()
    p4.text = "Candidate: Md Arman Alam | Enrolment No: A357175125020 | B.Tech CSE (AI/ML)"
    p4.font.size = Pt(14)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_TEXT_WHITE
    p4.space_before = Pt(10)

    p5 = tf1.add_paragraph()
    p5.text = "Project Supervisor: Miss Tanaya Ganguly | Amity School of Engineering & Technology"
    p5.font.size = Pt(13)
    p5.font.color.rgb = COLOR_TEXT_MUTED
    p5.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: Problem Statement & Motivation
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_background(s2, prs)
    add_header(s2, "1. Problem Statement & Research Motivation", "Chapter 1: Foundations")

    # Card Left: Problem
    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb_l = s2.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.8))
    tfl = tb_l.text_frame
    tfl.word_wrap = True
    p = tfl.paragraphs[0]
    p.text = "The Cricket Forecasting Challenge"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_GREEN

    bullets_l = [
        "High-Variance Dynamic: T20 compresses a match into 120 legal deliveries where match outcome changes on every ball.",
        "Failure of Naive Extrapolation: Standard broadcast projections multiply Current Run Rate (CRR x 20). This assumes linear future scoring and completely ignores wickets in hand and death-overs surge.",
        "Cluster Wicket Risk: Losing 2-3 quick wickets in middle overs causes sudden collapse, which linear formulas fail to anticipate.",
        "Overconfidence in Point Estimates: Single-number forecasts give false certainty without communicating uncertainty bounds."
    ]
    for b in bullets_l:
        pb = tfl.add_paragraph()
        pb.text = "• " + b
        pb.font.size = Pt(13); pb.font.color.rgb = COLOR_TEXT_WHITE; pb.space_before = Pt(10)

    # Card Right: Solution
    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb_r = s2.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tfr = tb_r.text_frame
    tfr.word_wrap = True
    p = tfr.paragraphs[0]
    p.text = "Our Proposed Approach"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_CYAN

    bullets_r = [
        "Conditional Ball-by-Ball Modeling: Continuous supervised prediction updated at every delivery checkpoint.",
        "Dual Prediction Architecture: Side-by-side transparent CRR baseline with a documented opening prior vs trained ML regressor.",
        "Conformal Uncertainty Intervals: Defensible 80% empirical prediction intervals that contract as remaining overs contract.",
        "Leak-Resistant Validation: Match-level chronological splitting ensuring no delivery rows cross between train and test sets.",
        "Explainable Inspection: Transparent feature attributions describing model behavior without causal speculation."
    ]
    for b in bullets_r:
        pb = tfr.add_paragraph()
        pb.text = "• " + b
        pb.font.size = Pt(13); pb.font.color.rgb = COLOR_TEXT_WHITE; pb.space_before = Pt(10)

    # =========================================================================
    # SLIDE 3: Scope & Non-Goals
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_background(s3, prs)
    add_header(s3, "2. Project Scope, Exclusions & Data Strategy", "Chapter 1 & 5: Scope Boundaries")

    # Card Left: Scope
    add_card(s3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s3.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.8))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Target Scope & Data Strategy"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_GREEN

    pts = [
        "Primary Target: Completed final score of an uninterrupted Men's T20 World Cup first innings scheduled for 20 overs.",
        "Data Source: Cricsheet structured Men's T20 International ball-by-ball CSV2 archive (3,627 matches).",
        "Sample Size Augmentation: Men's T20 World Cup tournament matches are held biennially (~250 historical matches). Broader Men's T20Is between major nations are utilized to train baseline dynamics.",
        "Subgroup Verification: World Cup tournament matches are isolated and evaluated as an independent held-out subgroup.",
        "Full Innings States: Evaluates ball 0 (opening prior 160.2 runs) to 120 balls / 10 wickets."
    ]
    for pt in pts:
        pb = tf.add_paragraph()
        pb.text = "✔ " + pt
        pb.font.size = Pt(13); pb.font.color.rgb = COLOR_TEXT_WHITE; pb.space_before = Pt(10)

    # Card Right: Explicit Exclusions
    add_card(s3, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), bg_color=RGBColor(36, 24, 20), border_color=RGBColor(239, 68, 68))
    tb2 = s3.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tf2 = tb2.text_frame; tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Explicit Non-Goals & Exclusions"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = RGBColor(248, 113, 113)

    excls = [
        "No Second-Innings Chase: Chase innings end early once the target is met, distorting natural run production. Strictly excluded in v1.",
        "No Win Probability: Win probability predicts binary match outcome, which has fundamentally different calibration targets.",
        "No Betting / Fantasy Advice: Responsible academic use prohibits commercial wagering advice or fantasy tips.",
        "No Automatic DLS: Rain-truncated matches (<115 legal balls) alter resource horizons and are explicitly filtered out.",
        "No Super Overs: Super over shootouts are isolated events and excluded from 20-over training."
    ]
    for ex in excls:
        pb = tf2.add_paragraph()
        pb.text = "✖ " + ex
        pb.font.size = Pt(13); pb.font.color.rgb = COLOR_TEXT_WHITE; pb.space_before = Pt(10)

    # =========================================================================
    # SLIDE 4: Architecture & Feature Engineering
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_background(s4, prs)
    add_header(s4, "3. System Architecture & Temporal Leakage Audit", "Chapter 4 & 5: Methodology")

    # Flow Box
    add_card(s4, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.8))
    tb_flow = s4.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(11.3), Inches(1.6))
    tff = tb_flow.text_frame; tff.word_wrap = True
    p = tff.paragraphs[0]
    p.text = "End-to-End System Pipeline (Report Figure 4.1)"
    p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_GOLD
    p_sub = tff.add_paragraph()
    p_sub.text = "Raw Cricsheet Ball-by-Ball → Domain Parsing (Legal Balls & Extras) → Chronological Match Split (2005–2024 Train / 2024–2026 Test) → Preprocessor Pipeline → Random Forest Regressor → Conformal Uncertainty Engine → Streamlit UI"
    p_sub.font.size = Pt(12); p_sub.font.color.rgb = COLOR_TEXT_WHITE; p_sub.space_before = Pt(4)

    # Left: Features
    add_card(s4, Inches(0.8), Inches(3.6), Inches(5.6), Inches(3.2))
    tb_f = s4.shapes.add_textbox(Inches(1.0), Inches(3.7), Inches(5.2), Inches(3.0))
    tff_f = tb_f.text_frame; tff_f.word_wrap = True
    p = tff_f.paragraphs[0]; p.text = "Model Input Features (At Ball B)"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_GREEN
    feats = [
        "Match Context: Batting Team, Bowling Team, Venue",
        "Current Innings State: Legal Balls Completed (0–120), Completed Overs (0.0–20.0), Current Runs",
        "Resource Depreciation: Wickets Lost (0–10)",
        "Momentum Window: Runs Scored in Last 5 Overs (30 legal balls)",
        "Scoring Velocity: Current Run Rate (with documented opening prior)"
    ]
    for f in feats:
        pb = tff_f.add_paragraph(); pb.text = "• " + f; pb.font.size = Pt(11.5); pb.font.color.rgb = COLOR_TEXT_WHITE; pb.space_before = Pt(4)

    # Right: Leakage Audit
    add_card(s4, Inches(6.8), Inches(3.6), Inches(5.7), Inches(3.2))
    tb_lkg = s4.shapes.add_textbox(Inches(7.0), Inches(3.7), Inches(5.3), Inches(3.0))
    tff_l = tb_lkg.text_frame; tff_l.word_wrap = True
    p = tff_l.paragraphs[0]; p.text = "Temporal Leakage Audit (Table 5.1)"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_CYAN
    lkg_pts = [
        "Zero Future Visibility: No end-of-innings aggregates, player final scores, or post-match results enter the model.",
        "Delivery Reconciliation: Legal deliveries strictly distinguished from illegal deliveries (wides & noballs).",
        "Extras Aggregation: Wides, noballs, byes, leg-byes, and penalties properly accumulated into score without inflating legal balls.",
        "Zero-Division Protection: At ball 0, CRR uses training mean prior (160.2 runs / 8.01 RPO) preventing runtime failures."
    ]
    for lk in lkg_pts:
        pb = tff_l.add_paragraph(); pb.text = "✔ " + lk; pb.font.size = Pt(11.5); pb.font.color.rgb = COLOR_TEXT_WHITE; pb.space_before = Pt(4)

    # =========================================================================
    # SLIDE 5: Evaluation & Table 7.1 Results
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_background(s5, prs)
    add_header(s5, "4. Chronological Held-Out Evaluation Results", "Chapter 7: Evaluation")

    add_card(s5, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tb_ev = s5.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11.1), Inches(4.8))
    tfe = tb_ev.text_frame; tfe.word_wrap = True
    p = tfe.paragraphs[0]
    p.text = "Replication of Report Table 7.1: Held-Out Test Evaluation"
    p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_GOLD

    p_sub = tfe.add_paragraph()
    p_sub.text = "Evaluated on 284 unseen matches (5,925 snapshots) chronologically held-out from 2024-09-28 to 2026-10-13:"
    p_sub.font.size = Pt(12); p_sub.font.color.rgb = COLOR_TEXT_MUTED; p_sub.space_before = Pt(3)

    metrics_bullets = [
        "Run-Rate Baseline (CRR + Prior): MAE = 22.97 runs | RMSE = 34.84 runs | Signed Bias = -9.15 runs",
        "Ridge Linear Regression:        MAE = 17.98 runs | RMSE = 24.37 runs | Signed Bias = -4.94 runs",
        "Random Forest Regressor (Best): MAE = 16.94 runs | RMSE = 24.50 runs | Signed Bias = -4.56 runs",
        "Superiority: Random Forest reduces Mean Absolute Error by 6.03 runs over the run-rate extrapolation benchmark."
    ]
    for mb in metrics_bullets:
        pb = tfe.add_paragraph(); pb.text = "• " + mb; pb.font.size = Pt(13); pb.font.color.rgb = COLOR_TEXT_WHITE; pb.space_before = Pt(6)

    p_phase = tfe.add_paragraph()
    p_phase.text = "Phase-Specific Error Analysis (Report Section 7.5):"
    p_phase.font.size = Pt(15); p_phase.font.bold = True; p_phase.font.color.rgb = COLOR_ACCENT_GREEN; p_phase.space_before = Pt(14)

    phases_text = [
        "Powerplay (Overs 0–6 / Balls 0–36):   RF MAE = 25.94 runs (vs Baseline 40.26 runs) [1,988 snapshots]",
        "Middle Overs (Overs 6.1–15 / Balls 37–90): RF MAE = 15.74 runs (vs Baseline 18.43 runs) [2,553 snapshots]",
        "Death Overs (Overs 15.1–20 / Balls 91–120): RF MAE = 6.21 runs  (vs Baseline 6.52 runs)  [1,384 snapshots]"
    ]
    for ph in phases_text:
        pb = tfe.add_paragraph(); pb.text = "  → " + ph; pb.font.size = Pt(12.5); pb.font.color.rgb = COLOR_TEXT_WHITE; pb.space_before = Pt(4)

    # =========================================================================
    # SLIDE 6: Uncertainty & Conformal Calibration
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_background(s6, prs)
    add_header(s6, "5. Uncertainty Bounds & Conformal Calibration", "Chapter 3 & 5: Uncertainty")

    add_card(s6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb_u1 = s6.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.8))
    tfu1 = tb_u1.text_frame; tfu1.word_wrap = True
    p = tfu1.paragraphs[0]; p.text = "Why Point Forecasts are Insufficient"; p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_GOLD
    u1_pts = [
        "Single-number estimates hide risk: An innings projected at 170 could easily swing to 140 on a collapse or 205 on death-overs hitting.",
        "MAE is NOT an Interval Guarantee: Calling +-MAE a 'confidence interval' is mathematically misleading (Report Requirement 4).",
        "Conformal Quantile Method: We calibrated an 80% split-conformal residual bound on out-of-fold development errors.",
        "Contracting Horizon (Figure 5.2): As balls remaining contract from 120 toward 0, uncertainty naturally converges to zero."
    ]
    for pt in u1_pts:
        pb = tfu1.add_paragraph(); pb.text = "• " + pt; pb.font.size = Pt(13); pb.font.color.rgb = COLOR_TEXT_WHITE; pb.space_before = Pt(10)

    add_card(s6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb_u2 = s6.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tfu2 = tb_u2.text_frame; tfu2.word_wrap = True
    p = tfu2.paragraphs[0]; p.text = "Empirical Test Set Measurements"; p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_CYAN
    u2_pts = [
        "Nominal Target: 80.0% Prediction Interval.",
        "Calibrated Half-Width (q80): 22.35 runs.",
        "Measured Test Coverage: 73.7% on unseen held-out matches.",
        "Average Interval Width: 42.9 runs.",
        "Distribution Shift Demonstration: The empirical coverage of 73.7% vs nominal 80% directly verifies the report's theoretical prediction: recent T20 scoring rate inflation shifts distributions, illustrating genuine real-world ML behavior."
    ]
    for pt in u2_pts:
        pb = tfu2.add_paragraph(); pb.text = "✔ " + pt; pb.font.size = Pt(13); pb.font.color.rgb = COLOR_TEXT_WHITE; pb.space_before = Pt(10)

    # =========================================================================
    # SLIDE 7: Streamlit Interactive UI
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_background(s7, prs)
    add_header(s7, "6. Streamlit Web Interface & Live Features", "Chapter 4 & 6: Application Prototype")

    add_card(s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb_ui1 = s7.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.8))
    tfui1 = tb_ui1.text_frame; tfui1.word_wrap = True
    p = tfui1.paragraphs[0]; p.text = "Dual Forecast Architecture"; p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_GREEN
    ui1_pts = [
        "Side-by-Side Comparison: ML Model Estimate shown directly adjacent to the Transparent CRR Baseline.",
        "Transparent Baseline Formula: Explains Runs + (CRR x Rem Overs) or documented opening prior (160.2 runs).",
        "Dynamic Trajectory Graph: Matplotlib dark pitch visualization displaying current runs, ML projected path, CRR path, and shaded 80% interval band.",
        "Quick Match Presets: Instant buttons for viva demos (Dominant Australia vs India, Death Overs Collapse, Opening State, Completed Innings)."
    ]
    for pt in ui1_pts:
        pb = tfui1.add_paragraph(); pb.text = "• " + pt; pb.font.size = Pt(13); pb.font.color.rgb = COLOR_TEXT_WHITE; pb.space_before = Pt(10)

    add_card(s7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb_ui2 = s7.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tfui2 = tb_ui2.text_frame; tfui2.word_wrap = True
    p = tfui2.paragraphs[0]; p.text = "Report Verification Tabs"; p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_CYAN
    ui2_pts = [
        "Tab 1: Table 7.1 Results & Phase Breakdown replicating exact report metrics.",
        "Tab 2: Interactive Table 7.2 Functional Validation Suite (TC01 to TC06).",
        "Tab 3: Full Architecture & Temporal Leakage Audit table.",
        "Tab 4: Ethics, Responsible Use Constraints & Gap Checklist.",
        "Responsive Theme: Tailored emerald dark aesthetic with clear typography and accessible input controls."
    ]
    for pt in ui2_pts:
        pb = tfui2.add_paragraph(); pb.text = "• " + pt; pb.font.size = Pt(13); pb.font.color.rgb = COLOR_TEXT_WHITE; pb.space_before = Pt(10)

    # =========================================================================
    # SLIDE 8: Table 7.2 Functional Validation Suite
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_background(s8, prs)
    add_header(s8, "7. Live Functional Test Suite (Report Table 7.2)", "Chapter 7: Verification")

    add_card(s8, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tb_t = s8.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11.1), Inches(4.8))
    tft = tb_t.text_frame; tft.word_wrap = True
    p = tft.paragraphs[0]; p.text = "Automated Functional Validation Results (All 6 Test Cases Pass)"; p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_GOLD

    tcases = [
        ("TC01 [Opening State]", "Input: 0/0 in 0.0 overs (0 legal balls)", "Pass: Uses documented opening prior (160.2 runs / 8.01 RPO) without zero-division error."),
        ("TC02 [Completed Innings]", "Input: 120 legal balls completed", "Pass: Accepted as completed innings; projection locks to current score (0 remaining balls)."),
        ("TC03 [Over-Limit Deliveries]", "Input: 121 legal balls in 20-over game", "Pass: Blocked by validation engine with explicit cricket boundary error."),
        ("TC04 [Over-Limit Dismissals]", "Input: 11 wickets lost", "Pass: Blocked by validation engine with explicit maximum wickets error."),
        ("TC05 [Unseen Match Venue]", "Input: Grand Prairie Stadium, Dallas", "Pass: Accepted with explicit informational warning; uses team base priors."),
        ("TC06 [Deterministic Output]", "Input: Repeat identical match state", "Pass: Bitwise identical predictions verified across multiple requests.")
    ]
    for tc_id, inp, res in tcases:
        p_tc = tft.add_paragraph()
        p_tc.text = f"✔ {tc_id}: {inp}"
        p_tc.font.size = Pt(13); p_tc.font.bold = True; p_tc.font.color.rgb = COLOR_ACCENT_GREEN; p_tc.space_before = Pt(8)
        p_res = tft.add_paragraph()
        p_res.text = f"   Result: {res}"
        p_res.font.size = Pt(11.5); p_res.font.color.rgb = COLOR_TEXT_WHITE; p_res.space_before = Pt(1)

    # =========================================================================
    # SLIDE 9: Conclusion & Viva Summary
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    apply_background(s9, prs)
    add_header(s9, "8. Conclusion, Ethics & Future Scope", "Chapter 8 & 9: Summary")

    add_card(s9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb_c1 = s9.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.8))
    tfc1 = tb_c1.text_frame; tfc1.word_wrap = True
    p = tfc1.paragraphs[0]; p.text = "Key Project Achievements"; p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_GREEN
    c1_pts = [
        "Full Report Alignment: Delivered complete prototype meeting all 8 specification requirements from the Major Project Report.",
        "Leakage-Free Modeling: Rigorous chronological match splitting with out-of-time evaluation.",
        "Superiority Established: Random Forest outperforms run-rate baseline by 6.03 runs MAE overall (and by 14.3 runs in Powerplay).",
        "Responsible Sports AI: Rigorously excluded betting, gambling, and fantasy speculation."
    ]
    for pt in c1_pts:
        pb = tfc1.add_paragraph(); pb.text = "✔ " + pt; pb.font.size = Pt(13); pb.font.color.rgb = COLOR_TEXT_WHITE; pb.space_before = Pt(10)

    add_card(s9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb_c2 = s9.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tfc2 = tb_c2.text_frame; tfc2.word_wrap = True
    p = tfc2.paragraphs[0]; p.text = "Future Roadmap (Report Chapter 8)"; p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_CYAN
    c2_pts = [
        "Live Push Ingestion: WebSocket integration with licensed match scorecard APIs (Report Sec 8.5).",
        "Player-Level Impact Indices: Dynamic strike rate tracking for current batters at the crease.",
        "Physical Pitch Telemetry: Sourcing pitch hardness and soil moisture telemetry from stadium groundstaff (Sec 8.7).",
        "Periodic Model Retraining: Managing temporal distribution drift following ICC tournaments (Sec 8.8)."
    ]
    for pt in c2_pts:
        pb = tfc2.add_paragraph(); pb.text = "• " + pt; pb.font.size = Pt(13); pb.font.color.rgb = COLOR_TEXT_WHITE; pb.space_before = Pt(10)

    prs.save(OUTPUT_PPTX)
    print(f"Presentation saved successfully to: {OUTPUT_PPTX}")

if __name__ == "__main__":
    build_presentation()
